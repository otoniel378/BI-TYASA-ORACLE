"""
mercado_fastmarkets/parser.py — Parseo tolerante de exports .xlsx de
Fastmarkets. Confirmado sobre 38 archivos reales que NO hay un formato
único: hay 2 layouts (serie de tiempo transpuesta, y tabla snapshot), y
dentro del layout de serie el orden de las filas de encabezado varía
(algunos archivos meten Currency/Unit Of Measure/Lot Size/Expiry Date/
Prompt entre Description y Symbol, y a veces Description va antes de
Symbol y a veces después). Por eso el parser busca cada etiqueta por su
valor de celda, no por posición fija de fila/columna.

Layout "serie": filas de encabezado con 'Symbol'/'Description' seguidas de
una fila con 'Date' (o 'Assessment Date') en columna A y 'Mid'/'Bid'/'Ask'
en las columnas de datos; debajo, una fila por fecha. Puede tener 1 o
varios símbolos compartiendo la columna de fecha (columna 0).

Layout "snapshot": una fila de encabezado normal (con 'Description' y
'Symbol' como nombres de columna, en cualquier posición) y una fila por
símbolo con su lectura más reciente — no es serie histórica.

Fastmarkets mezcla en la MISMA columna de fecha valores tipo datetime real
(fechas de pronóstico futuro generadas por su modelo) y texto 'd/m/Y'
(evaluaciones históricas reales) — confirmado en WIRE ROD (MB-STE-0192),
que trae historial hasta 2016 y pronóstico hasta 2028 en la misma columna.
TIPO se decide comparando la fecha contra "hoy" al momento de cargar, no
contra el tipo de celda (que es solo una pista indirecta).

Unidades: Fastmarkets exporta algunos widgets ya convertidos a $/tonne (nota
"Unit conversion has been applied..." en las primeras filas) y otros en su
unidad nativa sin convertir. Confirmado con datos reales que el MISMO símbolo
puede venir de dos archivos distintos con estados de conversión distintos
(MB-STE-0170 Rebar: un archivo ya convertido a tonne, ~1047 $/t, y otro sin
convertir en $/cwt, ~47 $/cwt — un salto de escala de ~22x justo en el corte
histórico/pronóstico si se mezclan tal cual). Por eso el parser normaliza a
tonne cuando detecta "$/cwt" en la Description y el archivo NO tiene la nota
de conversión — así el símbolo queda unificado sin importar de qué archivo
vino cada fila.
"""

import os
from datetime import date, datetime

import openpyxl

ETIQUETAS_FECHA = ("Date", "Assessment Date")
MEDIDAS_VALIDAS = ("Mid", "Bid", "Ask")
NOTA_CONVERSION = "Unit conversion has been applied"
FACTOR_CWT_A_TONELADA = 22.0462  # 1 tonelada métrica = 22.0462 cwt corto (EUA, 100 lb)


def _valor_texto(v) -> str:
    return str(v).strip() if v is not None else ""


def _parsear_fecha(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and v.strip():
        try:
            return datetime.strptime(v.strip(), "%d/%m/%Y").date()
        except ValueError:
            return None
    return None


def _parsear_numero(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def clasificar_hoja(filas: list) -> tuple:
    """Devuelve (tipo, índice_de_fila_clave) o (None, -1). Revisa las
    primeras 15 filas nada más — el encabezado siempre vive ahí."""
    limite = min(15, len(filas))
    for i in range(limite):
        fila = filas[i]
        if not fila:
            continue
        valores = [_valor_texto(c) for c in fila]
        c0 = valores[0] if valores else ""
        if c0 in ETIQUETAS_FECHA:
            return ("serie", i)
        if c0 != "Symbol" and "Symbol" in valores and "Description" in valores:
            return ("snapshot", i)
    return (None, -1)


def _tiene_nota_conversion(filas: list) -> bool:
    """True si el archivo ya declara haber convertido sus valores a
    tonelada — revisa solo las primeras filas, donde vive esa nota."""
    for fila in filas[:10]:
        if fila and any(c and NOTA_CONVERSION in str(c) for c in fila):
            return True
    return False


def _factor_normalizacion(descripcion: str, nota_conversion: bool) -> float:
    """1.0 salvo que la Description declare '$/cwt' y el archivo NO haya
    sido convertido por Fastmarkets — ahí se normaliza a $/tonne a mano
    para que el símbolo no quede con un salto de escala de ~22x según de
    qué archivo vino cada fila."""
    if not nota_conversion and descripcion and descripcion.rstrip().endswith("$/cwt"):
        return FACTOR_CWT_A_TONELADA
    return 1.0


def _buscar_fila_por_etiqueta(filas: list, idx_hasta: int, etiqueta: str):
    """Busca hacia atrás desde idx_hasta (sin incluir) una fila cuya
    columna A sea <etiqueta> — el orden real de filas varía entre archivos."""
    for i in range(idx_hasta - 1, -1, -1):
        fila = filas[i]
        if fila and _valor_texto(fila[0]) == etiqueta:
            return fila
    return None


def _valores_por_columna(fila) -> dict:
    """Mapa columna->texto para las celdas no vacías de una fila de
    encabezado (Symbol o Description), sin contar la columna 0 (que es la
    propia etiqueta 'Symbol'/'Description', no un dato)."""
    if not fila:
        return {}
    return {col: _valor_texto(v) for col, v in enumerate(fila) if col > 0 and _valor_texto(v)}


def parsear_serie(filas: list, idx_fecha: int, titulo: str | None, archivo: str, hoy: date,
                   nota_conversion: bool = False) -> list:
    fila_medida = filas[idx_fecha]
    fila_simbolo = _buscar_fila_por_etiqueta(filas, idx_fecha, "Symbol")
    fila_descripcion = _buscar_fila_por_etiqueta(filas, idx_fecha, "Description")

    simbolos_por_col = _valores_por_columna(fila_simbolo)
    descripciones_por_col = _valores_por_columna(fila_descripcion)

    # Archivos de un solo símbolo con varias columnas de medida (Low/Mid/
    # High/Mid Change %) solo repiten el símbolo/descripción en la primera
    # columna de datos (columna 1), no en cada columna de medida — se usa
    # como fallback para las demás. No se puede usar "único valor no vacío
    # en toda la fila" como criterio: algunos archivos meten una anotación
    # de footnote ("[P] Preliminary prices") en otra columna que coincide
    # por coincidencia con una columna de datos real, y eso rompería la
    # unicidad sin ser un símbolo/descripción de verdad.
    simbolo_fallback = simbolos_por_col.get(1)
    descripcion_fallback = descripciones_por_col.get(1)

    columnas = []
    for col in range(1, len(fila_medida)):
        medida = _valor_texto(fila_medida[col])
        if medida not in MEDIDAS_VALIDAS:
            continue
        symbol = simbolos_por_col.get(col) or simbolo_fallback or ""
        if not symbol:
            continue
        descripcion = descripciones_por_col.get(col) or descripcion_fallback or titulo or ""
        factor = _factor_normalizacion(descripcion, nota_conversion)
        columnas.append((col, symbol, descripcion, medida, factor))

    if not columnas:
        return []

    filas_out = []
    for fila in filas[idx_fecha + 1:]:
        if not fila or fila[0] is None:
            continue
        fecha = _parsear_fecha(fila[0])
        if fecha is None:
            continue
        tipo = "FORECAST" if fecha > hoy else "HISTORICO"
        for col, symbol, descripcion, medida, factor in columnas:
            if col >= len(fila):
                continue
            valor = _parsear_numero(fila[col])
            if valor is None:
                continue
            filas_out.append({
                "symbol": symbol[:50], "descripcion": descripcion[:500],
                "titulo_widget": (titulo or "")[:200], "fecha": fecha,
                "medida": medida, "valor": valor * factor, "tipo": tipo,
                "archivo_origen": archivo[:200],
            })
    return filas_out


def parsear_snapshot(filas: list, idx_encabezado: int, titulo: str | None, archivo: str, hoy: date,
                      nota_conversion: bool = False) -> list:
    encabezado = [_valor_texto(c) for c in filas[idx_encabezado]]
    try:
        col_desc = encabezado.index("Description")
        col_symbol = encabezado.index("Symbol")
        col_mid = encabezado.index("Mid")
    except ValueError:
        return []
    col_fecha = None
    for etiqueta in ETIQUETAS_FECHA:
        if etiqueta in encabezado:
            col_fecha = encabezado.index(etiqueta)
            break
    if col_fecha is None:
        return []

    filas_out = []
    for fila in filas[idx_encabezado + 1:]:
        # No filtrar por fila[0]: en algunos archivos (ej. "SCRAP") la
        # columna 0 es "Pricing rationale", casi siempre vacía, mientras
        # que Symbol/Description viven en otras columnas.
        if not fila:
            continue
        symbol = _valor_texto(fila[col_symbol]) if col_symbol < len(fila) else ""
        if not symbol:
            continue
        fecha = _parsear_fecha(fila[col_fecha]) if col_fecha < len(fila) else None
        if fecha is None:
            continue
        valor = _parsear_numero(fila[col_mid]) if col_mid < len(fila) else None
        if valor is None:
            continue
        descripcion = (_valor_texto(fila[col_desc]) if col_desc < len(fila) else "") or titulo or ""
        tipo = "FORECAST" if fecha > hoy else "HISTORICO"
        factor = _factor_normalizacion(descripcion, nota_conversion)
        filas_out.append({
            "symbol": symbol[:50], "descripcion": descripcion[:500],
            "titulo_widget": (titulo or "")[:200], "fecha": fecha,
            "medida": "Mid", "valor": valor * factor, "tipo": tipo,
            "archivo_origen": archivo[:200],
        })
    return filas_out


def parsear_archivo(ruta, hoy: date | None = None) -> tuple:
    """Devuelve (filas: list[dict], motivo_omision: str|None). No lanza
    excepción por un layout no reconocido — lo reporta para que el caller
    decida (revisar a mano, ampliar el parser, o descartar)."""
    hoy = hoy or date.today()
    nombre = os.path.basename(str(ruta))

    wb = openpyxl.load_workbook(str(ruta), read_only=True, data_only=True)
    try:
        ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.worksheets[0]
        filas = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    titulo = None
    for fila in filas[:5]:
        if fila and _valor_texto(fila[0]) == "Title":
            titulo = _valor_texto(fila[1]) if len(fila) > 1 else None
            break

    tipo_layout, idx = clasificar_hoja(filas)
    if tipo_layout is None:
        return [], "No se reconoció el layout (sin fila 'Symbol'/'Date'/'Assessment Date' en las primeras 15 filas)"

    nota_conversion = _tiene_nota_conversion(filas)
    if tipo_layout == "serie":
        resultado = parsear_serie(filas, idx, titulo, nombre, hoy, nota_conversion=nota_conversion)
    else:
        resultado = parsear_snapshot(filas, idx, titulo, nombre, hoy, nota_conversion=nota_conversion)

    if not resultado:
        return [], f"Layout '{tipo_layout}' reconocido pero no se extrajo ninguna fila válida"
    return resultado, None
